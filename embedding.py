import os, glob, requests, json
import numpy as np
from dotenv import load_dotenv

load_dotenv()

OPENROUTER_API_KEY = os.getenv("OPENROUTER_API_KEY")
if not OPENROUTER_API_KEY:
    raise RuntimeError("OPENROUTER_API_KEY missing")

EMBED_MODEL = "nvidia/llama-nemotron-embed-vl-1b-v2:free"


def parse_md(path):
    with open(path, "r", encoding="utf-8") as f:
        return f.read()

def parse_pdf(path):
    import fitz  # PyMuPDF
    doc = fitz.open(path)
    pages = []
    for page in doc:
        text = page.get_text().strip()
        if text:
            pages.append(text)
    doc.close()
    return "\n\n".join(pages)

def parse_docx(path):
    from docx import Document
    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)

def parse_txt(path):
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()

def parse_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = []
    for sheet in wb.worksheets:
        rows.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            row_text = " | ".join(str(cell) for cell in row if cell is not None)
            if row_text.strip():
                rows.append(row_text)
    return "\n".join(rows)

# Extension → parser mapping
PARSERS = {
    ".md":   parse_md,
    ".pdf":  parse_pdf,
    ".docx": parse_docx,
    ".txt":  parse_txt,
    ".xlsx": parse_xlsx,
}


def chunk_text(text, chunk_size=250, overlap=50):
    words = text.split()
    step = chunk_size - overlap
    out = []
    for i in range(0, len(words), step):
        chunk = " ".join(words[i:i + chunk_size])
        if chunk:
            out.append(chunk)
    return out


def load_and_chunk_docs(folder="data"):
    all_chunks = []
    supported_exts = set(PARSERS.keys())

    for path in glob.glob(os.path.join(folder, "**", "*"), recursive=True):
        ext = os.path.splitext(path)[1].lower()
        if ext not in supported_exts:
            continue

        print(f"  Loading [{ext}]: {path}")
        try:
            text = PARSERS[ext](path)
        except Exception as e:
            print(f"  ⚠ Skipping {path}: {e}")
            continue

        if not text.strip():
            print(f"  ⚠ Empty content, skipping: {path}")
            continue

        chunks = chunk_text(text, chunk_size=250, overlap=50)
        for idx, c in enumerate(chunks):
            all_chunks.append({
                "doc": path,
                "ext": ext,
                "chunk_id": idx,
                "text": c
            })

    return all_chunks



def embed_texts(texts, batch_size=32):
    url = "https://openrouter.ai/api/v1/embeddings"
    headers = {
        "Authorization": f"Bearer {OPENROUTER_API_KEY}",
        "Content-Type": "application/json",
    }
    all_embeddings = []
    for i in range(0, len(texts), batch_size):
        batch = texts[i:i + batch_size]
        data = {"model": EMBED_MODEL, "input": batch}
        r = requests.post(url, headers=headers, json=data, timeout=90)
        r.raise_for_status()
        all_embeddings.extend(item["embedding"] for item in r.json()["data"])
        print(f"  Embedded {min(i + batch_size, len(texts))}/{len(texts)} chunks")
    return all_embeddings



def main():
    chunks = load_and_chunk_docs("data")

    if not chunks:
        print("No documents found")
        return

    texts = [c["text"] for c in chunks]
    embeddings = embed_texts(texts)

    with open("data/chunks.json", "w", encoding="utf-8") as f:
        json.dump(chunks, f, ensure_ascii=False, indent=2)

    with open("data/embeddings.npy", "wb") as f:
        np.save(f, np.array(embeddings, dtype="float32"))


if __name__ == "__main__":
    main()